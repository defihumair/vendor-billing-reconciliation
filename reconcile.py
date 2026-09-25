import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REQUIRED = {
    "internal volumes": ["year", "week", "facility", "billing_head", "qty"],
    "vendor invoice": ["year", "week", "invoice_no", "facility", "billing_head", "qty", "amount"],
    "rate card": ["billing_head", "uom", "rate"],
}
OVERRIDES = ["System default", "Force internal volume", "Force vendor volume"]
STATUS_ORDER = {"Overbilled": 0, "Overbilled + rate": 0, "Underbilled": 1, "Underbilled + rate": 1, "Match + rate": 2, "Match": 3}
TAX_RATE = 0.15
QTY_TOLERANCE = 0.5
RATE_TOLERANCE = 1.0


def _prepare(df, name):
    df = df.copy()
    df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
    missing = [c for c in REQUIRED[name] if c not in df.columns]
    if missing:
        raise ValueError(f"The {name} file is missing column(s): {', '.join(missing)}. Expected: {', '.join(REQUIRED[name])}.")
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].astype(str).str.strip()
    return df


def validate(internal, vendor, rates):
    i = _prepare(internal, "internal volumes")
    v = _prepare(vendor, "vendor invoice")
    r = _prepare(rates, "rate card")
    for df, cols in ((i, ["qty"]), (v, ["qty", "amount"]), (r, ["rate"])):
        for col in cols:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    for df in (i, v):
        for col in ("year", "week"):
            df[col] = pd.to_numeric(df[col], errors="coerce")
        df.dropna(subset=["year", "week"], inplace=True)
        df["year"] = df["year"].astype(int)
        df["week"] = df["week"].astype(int)
    unknown = sorted((set(i["billing_head"]) | set(v["billing_head"])) - set(r["billing_head"]))
    if unknown:
        raise ValueError(f"No contract rate found for: {', '.join(unknown)}. Add them to the rate card.")
    if v.empty:
        raise ValueError("The vendor invoice file has no rows.")
    return i, v, r


def periods(vendor):
    return sorted({(int(y), int(w)) for y, w in zip(vendor["year"], vendor["week"])}, reverse=True)


def invoice_number(vendor, year, week):
    s = vendor.loc[(vendor["year"] == year) & (vendor["week"] == week), "invoice_no"].dropna()
    return str(s.iloc[0]) if len(s) else f"INV-{year}-W{week:02d}"


def reconcile(internal, vendor, rates, year, week):
    iw = (internal[(internal["year"] == year) & (internal["week"] == week)]
          .groupby(["billing_head", "facility"], as_index=False)["qty"].sum()
          .rename(columns={"qty": "internal_qty"}))
    vw = (vendor[(vendor["year"] == year) & (vendor["week"] == week)]
          .groupby(["billing_head", "facility"], as_index=False)[["qty", "amount"]].sum()
          .rename(columns={"qty": "vendor_qty", "amount": "vendor_amount"}))
    df = iw.merge(vw, on=["billing_head", "facility"], how="outer")
    df[["internal_qty", "vendor_qty", "vendor_amount"]] = df[["internal_qty", "vendor_qty", "vendor_amount"]].fillna(0.0)
    df = df.merge(rates[["billing_head", "uom", "rate"]], on="billing_head", how="left")
    df = df[(df["internal_qty"] > 0) | (df["vendor_qty"] > 0) | (df["vendor_amount"] > 0)].copy()
    df["qty_variance"] = df["vendor_qty"] - df["internal_qty"]
    df["rate_gap"] = df["vendor_amount"] - df["vendor_qty"] * df["rate"]
    status, note, default_qty = [], [], []
    for row in df.itertuples(index=False):
        if row.qty_variance > QTY_TOLERANCE:
            s, t, q = "Overbilled", "Vendor billed more than internal records: pay internal volume", row.internal_qty
        elif row.qty_variance < -QTY_TOLERANCE:
            s, t, q = "Underbilled", "Vendor billed less than internal records: pay as invoiced", row.vendor_qty
        else:
            s, t, q = "Match", "Volumes match", row.vendor_qty
        if abs(row.rate_gap) > RATE_TOLERANCE:
            s += " + rate"
            t += "; vendor rate differs from contract, contract rate applied"
        status.append(s)
        note.append(t)
        default_qty.append(q)
    df["status"] = status
    df["recommendation"] = note
    df["default_qty"] = default_qty
    df["sort"] = df["status"].map(STATUS_ORDER)
    df = df.sort_values(["sort", "billing_head", "facility"]).drop(columns=["sort", "rate_gap"]).reset_index(drop=True)
    df.insert(0, "line", [f"L{n:02d}" for n in range(1, len(df) + 1)])
    df["override"] = "System default"
    return df.rename(columns={
        "line": "Line", "billing_head": "Billing head", "facility": "Facility", "uom": "UOM", "rate": "Rate",
        "internal_qty": "Internal qty", "vendor_qty": "Vendor qty", "vendor_amount": "Vendor amount",
        "qty_variance": "Qty variance", "status": "Status", "recommendation": "Recommendation",
        "default_qty": "Default qty", "override": "Override",
    })


def apply_overrides(lines, overrides=None):
    df = lines.copy()
    if overrides is not None:
        df["Override"] = list(overrides)
    df["Override"] = df["Override"].where(df["Override"].isin(OVERRIDES), "System default")
    approved = df["Default qty"].copy()
    approved = approved.mask(df["Override"] == "Force internal volume", df["Internal qty"])
    approved = approved.mask(df["Override"] == "Force vendor volume", df["Vendor qty"])
    df["Approved qty"] = approved
    df["Payable"] = (df["Approved qty"] * df["Rate"]).round(2)
    df["Correction"] = (df["Vendor amount"] - df["Payable"]).round(2)
    return df


def summarize(final):
    internal_value = float((final["Internal qty"] * final["Rate"]).sum())
    vendor_total = float(final["Vendor amount"].sum())
    payable = float(final["Payable"].sum())
    return {
        "internal_value": round(internal_value, 2),
        "vendor_total": round(vendor_total, 2),
        "payable": round(payable, 2),
        "correction": round(vendor_total - payable, 2),
        "lines": int(len(final)),
        "exceptions": int((final["Status"] != "Match").sum()),
        "overrides": int((final["Override"] != "System default").sum()),
    }


def savings_log(internal, vendor, rates):
    rows = []
    for year, week in sorted(periods(vendor)):
        final = apply_overrides(reconcile(internal, vendor, rates, year, week))
        s = summarize(final)
        rows.append({
            "Year": year, "Week": week, "Invoice": invoice_number(vendor, year, week),
            "Vendor invoiced": s["vendor_total"], "Approved payable": s["payable"],
            "Correction": s["correction"], "Exception lines": s["exceptions"], "Lines": s["lines"],
        })
    return pd.DataFrame(rows)


def _style_header(cells):
    fill = PatternFill("solid", start_color="0F5E6B")
    for c in cells:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_invoice(final, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    thin = Side(style="thin", color="B7C4CC")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws["A1"] = "Payable invoice"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    details = [
        ("Vendor", meta["vendor"]), ("Bill to", meta["bill_to"]), ("Invoice no.", meta["invoice_no"]),
        ("Invoice date", meta["invoice_date"]), ("Billing week", f"{meta['year']} week {meta['week']}"),
        ("Currency", meta["currency"]), ("Approved by", meta["approver"]),
    ]
    for n, (k, v) in enumerate(details, start=3):
        ws.cell(n, 1, k).font = Font(name="Arial", bold=True)
        ws.cell(n, 2, v).font = Font(name="Arial")
    row = len(details) + 4
    subtotal_cells = []
    for facility in sorted(final["Facility"].unique()):
        block = final[(final["Facility"] == facility) & (final["Approved qty"] > 0)]
        if block.empty:
            continue
        ws.cell(row, 1, facility).font = Font(name="Arial", bold=True, size=11)
        row += 1
        for col, title in enumerate(["Billing head", "UOM", "Qty", "Rate", "Amount"], start=1):
            ws.cell(row, col, title)
        _style_header([ws.cell(row, c) for c in range(1, 6)])
        row += 1
        first = row
        for head, uom, qty, rate in block[["Billing head", "UOM", "Approved qty", "Rate"]].itertuples(index=False, name=None):
            ws.cell(row, 1, head)
            ws.cell(row, 2, uom)
            ws.cell(row, 3, float(qty)).number_format = "#,##0.00"
            ws.cell(row, 4, float(rate)).number_format = "#,##0.00"
            ws.cell(row, 5, f"=C{row}*D{row}").number_format = "#,##0.00"
            for c in range(1, 6):
                ws.cell(row, c).border = box
                ws.cell(row, c).font = Font(name="Arial")
            row += 1
        ws.cell(row, 4, "Subtotal").font = Font(name="Arial", bold=True)
        ws.cell(row, 5, f"=SUM(E{first}:E{row - 1})").number_format = "#,##0.00"
        ws.cell(row, 5).font = Font(name="Arial", bold=True)
        subtotal_cells.append(f"E{row}")
        row += 2
    total_row = row
    ws.cell(total_row, 4, "Total before tax").font = Font(name="Arial", bold=True)
    ws.cell(total_row, 5, "=" + ("+".join(subtotal_cells) if subtotal_cells else "0")).number_format = "#,##0.00"
    ws.cell(total_row + 1, 4, f"Sales tax ({int(TAX_RATE * 100)}%)").font = Font(name="Arial", bold=True)
    ws.cell(total_row + 1, 5, f"=E{total_row}*{TAX_RATE}").number_format = "#,##0.00"
    ws.cell(total_row + 2, 4, "Total payable").font = Font(name="Arial", bold=True)
    ws.cell(total_row + 2, 5, f"=E{total_row}+E{total_row + 1}").number_format = "#,##0.00"
    for r in range(total_row, total_row + 3):
        ws.cell(r, 5).font = Font(name="Arial", bold=True)
        ws.cell(r, 5).border = box
    for col, width in zip("ABCDE", (34, 12, 14, 12, 18)):
        ws.column_dimensions[col].width = width

    rs = wb.create_sheet("Reconciliation")
    cols = ["Line", "Billing head", "Facility", "UOM", "Rate", "Internal qty", "Vendor qty", "Vendor amount",
            "Qty variance", "Status", "Override", "Approved qty", "Payable", "Correction", "Recommendation"]
    for c, title in enumerate(cols, start=1):
        rs.cell(1, c, title)
    _style_header([rs.cell(1, c) for c in range(1, len(cols) + 1)])
    for r, rec in enumerate(final[cols].itertuples(index=False), start=2):
        for c, value in enumerate(rec, start=1):
            cell = rs.cell(r, c, value.item() if hasattr(value, "item") else value)
            cell.font = Font(name="Arial")
            if isinstance(cell.value, (int, float)) and c >= 5:
                cell.number_format = "#,##0.00"
    last = len(final) + 1
    rs.cell(last + 1, 7, "Totals").font = Font(name="Arial", bold=True)
    for c in (8, 13, 14):
        letter = get_column_letter(c)
        cell = rs.cell(last + 1, c, f"=SUM({letter}2:{letter}{last})")
        cell.font = Font(name="Arial", bold=True)
        cell.number_format = "#,##0.00"
    for c, width in enumerate([7, 30, 14, 9, 9, 12, 12, 15, 12, 18, 20, 13, 15, 13, 60], start=1):
        rs.column_dimensions[get_column_letter(c)].width = width
    rs.freeze_panes = "A2"

    au = wb.create_sheet("Audit trail")
    s = summarize(final)
    audit = [
        ("Approved by", meta["approver"]), ("Approved at", meta["approved_at"]), ("Invoice no.", meta["invoice_no"]),
        ("Billing week", f"{meta['year']} week {meta['week']}"), ("Lines reviewed", s["lines"]),
        ("Exception lines", s["exceptions"]), ("Manager overrides", s["overrides"]),
        ("Vendor invoiced", s["vendor_total"]), ("Approved payable (before tax)", s["payable"]),
        ("Correction", s["correction"]),
    ]
    for n, (k, v) in enumerate(audit, start=1):
        au.cell(n, 1, k).font = Font(name="Arial", bold=True)
        cell = au.cell(n, 2, v)
        cell.font = Font(name="Arial")
        if isinstance(v, float):
            cell.number_format = "#,##0.00"
    start = len(audit) + 2
    heads = ["Line", "Billing head", "Facility", "Default qty", "Override", "Approved qty"]
    for c, title in enumerate(heads, start=1):
        au.cell(start, c, title)
    _style_header([au.cell(start, c) for c in range(1, len(heads) + 1)])
    changed = final[final["Override"] != "System default"]
    if changed.empty:
        au.cell(start + 1, 1, "No manager overrides; system defaults applied to every line.").font = Font(name="Arial", italic=True)
    for r, rec in enumerate(changed[heads].itertuples(index=False), start=start + 1):
        for c, value in enumerate(rec, start=1):
            au.cell(r, c, value.item() if hasattr(value, "item") else value).font = Font(name="Arial")
    for col, width in zip("ABCDEF", (30, 30, 14, 13, 22, 13)):
        au.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def invoice_meta(vendor_name, bill_to, invoice_no, invoice_date, year, week, currency, approver):
    return {
        "vendor": vendor_name, "bill_to": bill_to, "invoice_no": invoice_no,
        "invoice_date": invoice_date.strftime("%d-%b-%Y") if hasattr(invoice_date, "strftime") else str(invoice_date),
        "year": year, "week": week, "currency": currency, "approver": approver,
        "approved_at": datetime.now().strftime("%d-%b-%Y %H:%M"),
    }
